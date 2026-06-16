"""导出：双语对照 HTML / Markdown / Word(docx)。"""
import html
import io
import re

from . import storage


def _esc(t: str) -> str:
    return html.escape(t or "")


def export_html(book_id: int) -> bytes:
    data = storage.book_full(book_id)
    if not data:
        return b""
    book = data["book"]
    parts = [
        "<!DOCTYPE html><html lang='zh'><head><meta charset='utf-8'>",
        f"<title>{_esc(book['title'])} · 双语对照</title>",
        """<style>
        body{font-family:-apple-system,'Segoe UI','Microsoft YaHei',sans-serif;max-width:1100px;
        margin:0 auto;padding:24px;color:#222;line-height:1.7}
        h1{text-align:center}
        .page{border-top:2px dashed #ddd;margin-top:32px;padding-top:12px}
        .pn{color:#999;font-size:13px}
        .row{display:flex;gap:24px;margin:10px 0}
        .col{flex:1}
        .en{color:#555}
        .zh{color:#111}
        .h .en,.h .zh{font-weight:700;font-size:1.15em}
        .note{background:#fff8e6;border-left:4px solid #f0b400;padding:10px 14px;margin:12px 0;
        border-radius:4px;font-size:14px;white-space:pre-wrap}
        .unote{background:#eef6ff;border-left:4px solid #3b82f6;padding:10px 14px;border-radius:4px;font-size:14px;white-space:pre-wrap}
        </style></head><body>""",
        f"<h1>{_esc(book['title'])}</h1>",
        "<p style='text-align:center;color:#999'>英中双语对照 · 由 Book Translator 生成</p>",
    ]
    for page in sorted(data["blocks_by_page"].keys()):
        parts.append(f"<div class='page'><div class='pn'>第 {page} 页</div>")
        for b in data["blocks_by_page"][page]:
            cls = "row h" if b["type"] == "h" else "row"
            parts.append(
                f"<div class='{cls}'><div class='col en'>{_esc(b['original'])}</div>"
                f"<div class='col zh'>{_esc(b['translation'])}</div></div>"
            )
        note = data["notes"].get(page, {})
        if note.get("ai_note"):
            parts.append(f"<div class='note'><b>AI 笔记</b><br>{_esc(note['ai_note'])}</div>")
        if note.get("user_note"):
            parts.append(f"<div class='unote'><b>我的笔记</b><br>{_esc(note['user_note'])}</div>")
        parts.append("</div>")
    parts.append("</body></html>")
    return "\n".join(parts).encode("utf-8")


def export_markdown(book_id: int) -> bytes:
    data = storage.book_full(book_id)
    if not data:
        return b""
    book = data["book"]
    out = [f"# {book['title']}\n", "> 英中双语对照 · 由 Book Translator 生成\n"]
    for page in sorted(data["blocks_by_page"].keys()):
        out.append(f"\n---\n\n### 第 {page} 页\n")
        for b in data["blocks_by_page"][page]:
            prefix = "#### " if b["type"] == "h" else ""
            out.append(f"{prefix}{b['original']}\n")
            out.append(f"{prefix}**{b['translation']}**\n")
        note = data["notes"].get(page, {})
        if note.get("ai_note"):
            out.append(f"\n> 📝 **AI 笔记**\n>\n> {note['ai_note'].replace(chr(10), chr(10)+'> ')}\n")
        if note.get("user_note"):
            out.append(f"\n> ✍️ **我的笔记**\n>\n> {note['user_note'].replace(chr(10), chr(10)+'> ')}\n")
    return "\n".join(out).encode("utf-8")


def export_docx(book_id: int) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    data = storage.book_full(book_id)
    if not data:
        return b""
    book = data["book"]
    doc = Document()
    doc.add_heading(book["title"] or "双语对照", level=0)
    for page in sorted(data["blocks_by_page"].keys()):
        doc.add_paragraph(f"第 {page} 页").italic = True
        for b in data["blocks_by_page"][page]:
            if b["type"] == "h":
                doc.add_heading(b["original"], level=2)
                doc.add_heading(b["translation"], level=2)
            else:
                p = doc.add_paragraph()
                r = p.add_run(b["original"])
                r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
                r.font.size = Pt(10)
                doc.add_paragraph(b["translation"])
        note = data["notes"].get(page, {})
        if note.get("ai_note"):
            p = doc.add_paragraph()
            p.add_run("【AI 笔记】" + note["ai_note"]).italic = True
        if note.get("user_note"):
            p = doc.add_paragraph()
            p.add_run("【我的笔记】" + note["user_note"]).italic = True
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def export_pdf(book_id: int) -> bytes:
    """生成带目录(可点击)和书签的双语对照 PDF。中文用内置 CID 字体，无需额外字体文件。"""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import (
        BaseDocTemplate, Frame, PageBreak, PageTemplate, Paragraph,
        Spacer, Table, TableStyle,
    )
    from reportlab.platypus.tableofcontents import TableOfContents

    data = storage.book_full(book_id)
    if not data:
        return b""
    book = data["book"]

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    ZH = "STSong-Light"   # 中文（也含 ASCII）
    EN = "Helvetica"      # 英文

    en_style = ParagraphStyle("en", fontName=EN, fontSize=9.5, leading=14, textColor=colors.HexColor("#555555"))
    zh_style = ParagraphStyle("zh", fontName=ZH, fontSize=10.5, leading=16, textColor=colors.black)
    en_h = ParagraphStyle("enh", fontName="Helvetica-Bold", fontSize=12, leading=16, textColor=colors.HexColor("#333333"))
    zh_h = ParagraphStyle("zhh", fontName=ZH, fontSize=13, leading=18, textColor=colors.black)
    note_style = ParagraphStyle("note", fontName=ZH, fontSize=9.5, leading=15,
                                textColor=colors.HexColor("#5a4a00"), backColor=colors.HexColor("#fff8e6"),
                                borderPadding=6, leftIndent=2, spaceBefore=4, spaceAfter=4)
    unote_style = ParagraphStyle("unote", fontName=ZH, fontSize=9.5, leading=15,
                                 textColor=colors.HexColor("#0b3d91"), backColor=colors.HexColor("#eef6ff"),
                                 borderPadding=6, spaceBefore=4, spaceAfter=4)
    pn_style = ParagraphStyle("pn", fontName=ZH, fontSize=8, textColor=colors.HexColor("#999999"), spaceBefore=10)
    title_style = ParagraphStyle("title", fontName=ZH, fontSize=22, leading=28, alignment=1, spaceAfter=8)
    sub_style = ParagraphStyle("sub", fontName=ZH, fontSize=11, alignment=1, textColor=colors.HexColor("#888888"))

    # 带书签/目录回调的文档模板
    class BookDoc(BaseDocTemplate):
        def __init__(self, *a, **kw):
            super().__init__(*a, **kw)
            self._heading_seq = 0

        def beforeDocument(self):
            # 每遍构建重置计数，保证书签 key 在多遍间稳定
            self._heading_seq = 0

        def afterFlowable(self, flowable):
            if isinstance(flowable, Paragraph) and getattr(flowable, "style", None):
                if flowable.style.name == "zhh":
                    text = flowable.getPlainText()
                    if not text:
                        return
                    key = "h%d" % self._heading_seq
                    self._heading_seq += 1
                    self.canv.bookmarkPage(key)
                    self.canv.addOutlineEntry(text, key, level=0, closed=False)
                    self.notify("TOCEntry", (0, text, self.page, key))

    buf = io.BytesIO()
    doc = BookDoc(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
                  topMargin=18 * mm, bottomMargin=16 * mm, title=book["title"])
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="main")
    doc.addPageTemplates([PageTemplate(id="main", frames=[frame])])

    def cell_table(left_para, right_para):
        col = (doc.width - 8) / 2.0
        t = Table([[left_para, right_para]], colWidths=[col, col])
        t.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 0),
            ("RIGHTPADDING", (0, 0), (0, 0), 8),
            ("RIGHTPADDING", (1, 0), (1, 0), 0),
            ("TOPPADDING", (0, 0), (-1, -1), 2),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            ("LINEBELOW", (0, 0), (-1, -1), 0.3, colors.HexColor("#eeeeee")),
        ]))
        return t

    story = []
    # 封面
    story.append(Spacer(1, 60 * mm))
    story.append(Paragraph(_esc(book["title"]), title_style))
    story.append(Paragraph("英中双语对照 · 由 Book Translator 生成", sub_style))
    story.append(PageBreak())

    # 目录
    toc = TableOfContents()
    toc.levelStyles = [ParagraphStyle("toc0", fontName=ZH, fontSize=11, leading=18,
                                      textColor=colors.black)]
    toc_title = ParagraphStyle("toctitle", fontName=ZH, fontSize=16, leading=22, textColor=colors.black)
    story.append(Paragraph("目录", toc_title))
    story.append(Spacer(1, 6 * mm))
    story.append(toc)
    story.append(PageBreak())

    # 正文
    for page in sorted(data["blocks_by_page"].keys()):
        story.append(Paragraph(f"— 第 {page} 页 —", pn_style))
        for b in data["blocks_by_page"][page]:
            if b["type"] == "h":
                # 标题整行排版，触发书签/目录；中文标题用 zhh 样式
                story.append(Paragraph(_esc(b["original"]), en_h))
                story.append(Paragraph(_esc(b["translation"] or b["original"]), zh_h))
            else:
                story.append(cell_table(
                    Paragraph(_esc(b["original"]), en_style),
                    Paragraph(_esc(b["translation"]), zh_style),
                ))
        note = data["notes"].get(page, {})
        if note.get("ai_note"):
            story.append(Paragraph("📝 AI 笔记<br/>" + _esc(note["ai_note"]).replace("\n", "<br/>"), note_style))
        if note.get("user_note"):
            story.append(Paragraph("✍️ 我的笔记<br/>" + _esc(note["user_note"]).replace("\n", "<br/>"), unote_style))

    doc.multiBuild(story)  # 两遍构建以解析目录页码
    return buf.getvalue()


def safe_filename(name: str) -> str:
    name = re.sub(r"[^\w一-龥\-. ]", "_", name or "book").strip()
    return name or "book"
